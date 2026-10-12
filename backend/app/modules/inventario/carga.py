"""
CARGA INICIAL (2026-10-12) — el restaurante arranca de cero.

El sistema genera una PLANTILLA (Excel con tres hojas, o un CSV por hoja):
Insumos, Platos y Recetas. El restaurante la llena, la sube, y antes de
guardar nada el sistema la revisa FILA POR FILA y devuelve un informe con
cada error (hoja, fila, columna y causa). Solo si no hay ningún error se
puede aplicar, y se aplica entera en UNA transacción: o entra todo o nada.

Qué hace cada hoja al aplicarse:
- Insumos: crea los que no existen (por nombre) y actualiza grupo, mínimo y
  costo de los que ya existen. La EXISTENCIA nunca se escribe directo: si la
  fila trae existencia, se registra un movimiento ("inicial" si el insumo es
  nuevo; "ajuste" por la diferencia si ya existía), con motivo "Carga inicial".
- Platos: crea los que no existen (por nombre) y actualiza precio,
  categoría, cocina y descripción de los existentes. Una categoría que no
  existe se crea.
- Recetas: la receta de cada plato que aparece en la hoja se REEMPLAZA por
  las filas de la hoja (un plato que no aparece conserva la suya).

Números: en celdas de Excel no hay ambigüedad. En texto (CSV) se acepta
coma decimal ("1,5") y punto de miles solo con grupos de tres ("16.000").
"""
import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.core.branding import NOMBRE_PRODUCTO
from app.core.error_handlers import traducir_error
from app.core.errors import AppError
from app.modules.catalogo import schemas as sc
from app.modules.catalogo.models import Categoria, Producto, RecetaItem
from app.modules.inventario import schemas as si, services
from app.modules.inventario.models import Insumo
from app.modules.inventario.unidades import compatibles, convertir

MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 5000
TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TIPOS_CSV = ("text/csv", "application/csv", "text/plain", "application/vnd.ms-excel")

# hoja → [(encabezado, clave, obligatoria, ejemplo, ayuda)]
COLUMNAS: dict[str, list[tuple[str, str, bool, list, str]]] = {
    "Insumos": [
        ("Nombre", "nombre", True, ["Arroz", "Queso mozzarella"], "Como lo compras"),
        ("Unidad", "unidad", True, ["kg", "g"], "kg, g, L, ml o unidad"),
        ("Grupo", "grupo", False, ["Abarrotes", "Lácteos"], "Opcional: Carnes, Verduras…"),
        ("Mínimo", "minimo", False, [5, 500], "Por debajo se avisa que hay que comprar"),
        ("Costo unitario", "costo_unitario", False, [4200, 28], "Costo de UNA unidad (un kg, un g…)"),
        ("Existencia", "existencia", False, [25, 3000], "Lo que hay hoy en la bodega"),
    ],
    "Platos": [
        ("Nombre", "nombre", True, ["Arroz con pollo", "Gaseosa 400 ml"], ""),
        ("Categoría", "categoria", False, ["Almuerzos", "Bebidas"], "Si no existe, se crea"),
        ("Precio", "precio", True, [14000, 3500], "Precio de venta"),
        ("Va a cocina", "va_a_cocina", False, ["Sí", "No"], "No = venta directa (no sale en la comanda)"),
        ("Código", "sku", False, ["", ""], "Opcional"),
        ("Descripción", "descripcion", False, ["Con ensalada", ""], "Opcional"),
    ],
    "Recetas": [
        ("Plato", "plato", True, ["Arroz con pollo", "Arroz con pollo"], "Igual que en la hoja Platos"),
        ("Insumo", "insumo", True, ["Arroz", "Queso mozzarella"], "Igual que en la hoja Insumos"),
        ("Cantidad", "cantidad", True, [200, 30], "Por porción"),
        ("Unidad", "unidad", True, ["g", "g"], "Puede ser distinta a la del insumo si es de la misma familia"),
        ("Merma %", "merma_porcentaje", False, [0, 0], "Opcional, de 0 a 90"),
    ],
}
HOJAS = tuple(COLUMNAS)


class ArchivoInvalido(AppError):
    status_code = 422
    detail = "El archivo no se pudo leer."


# ---------------------------------------------------------------------------
# Plantilla
# ---------------------------------------------------------------------------
INSTRUCCIONES = [
    "Plantilla de carga inicial",
    "",
    "1. Llena las hojas Insumos, Platos y Recetas (puedes dejar alguna vacía).",
    "2. No cambies los encabezados de la fila 1. Las filas de ejemplo se pueden borrar o reemplazar.",
    "3. Los nombres se comparan sin importar mayúsculas: si un insumo o plato ya existe, se actualiza.",
    "4. La existencia no se escribe a mano: queda como un movimiento 'Carga inicial' en el historial.",
    "5. Sube el archivo en Inventario → Carga inicial. Primero se revisa fila por fila y te decimos cada error;",
    "   nada se guarda hasta que el archivo esté limpio y confirmes.",
]


def plantilla_xlsx() -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Instrucciones"
    for i, linea in enumerate(INSTRUCCIONES, start=1):
        hoja.cell(row=i, column=1, value=linea)
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.column_dimensions["A"].width = 110
    gris = PatternFill("solid", fgColor="F1F1F4")
    for nombre, columnas in COLUMNAS.items():
        h = libro.create_sheet(nombre)
        for j, (encabezado, _, obligatoria, ejemplos, ayuda) in enumerate(columnas, start=1):
            celda = h.cell(row=1, column=j, value=f"{encabezado}{' *' if obligatoria else ''}")
            celda.font = Font(bold=True)
            celda.fill = gris
            h.column_dimensions[celda.column_letter].width = max(16, len(encabezado) + 6)
            for k, ejemplo in enumerate(ejemplos, start=2):
                if ejemplo != "":
                    h.cell(row=k, column=j, value=ejemplo)
            if ayuda:
                celda.comment = Comment(ayuda, NOMBRE_PRODUCTO)
        if nombre in ("Insumos", "Recetas"):
            col = "B" if nombre == "Insumos" else "D"
            lista = DataValidation(type="list", formula1='"kg,g,L,ml,unidad"', allow_blank=False)
            lista.error = "Usa kg, g, L, ml o unidad"
            h.add_data_validation(lista)
            lista.add(f"{col}2:{col}{MAX_FILAS}")
        if nombre == "Platos":
            si_no = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
            h.add_data_validation(si_no)
            si_no.add(f"D2:D{MAX_FILAS}")
        h.freeze_panes = "A2"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def plantilla_csv(hoja: str) -> bytes:
    columnas = COLUMNAS[hoja]
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    escritor.writerow([f"{c[0]}{' *' if c[2] else ''}" for c in columnas])
    for k in range(len(columnas[0][3])):
        escritor.writerow([str(c[3][k]).replace(".", ",") for c in columnas])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------
def _normalizar(texto) -> str:
    return re.sub(r"[\s*]+", " ", str(texto or "")).strip().lower()


ALIAS = {hoja: {_normalizar(c[0]): c[1] for c in cols} for hoja, cols in COLUMNAS.items()}


def _filas_de(encabezados: list, filas: list[list], hoja: str) -> tuple[list[dict], list[dict]]:
    """Convierte filas crudas en dicts {clave: valor, "_fila": n}. Errores de
    encabezado si falta una columna obligatoria."""
    indice = {}
    for j, encabezado in enumerate(encabezados):
        clave = ALIAS[hoja].get(_normalizar(encabezado))
        if clave:
            indice[clave] = j
    errores = []
    for encabezado, clave, obligatoria, _, _ in COLUMNAS[hoja]:
        if obligatoria and clave not in indice:
            errores.append({"hoja": hoja, "fila": 1, "columna": encabezado, "mensaje": f"Falta la columna «{encabezado}» en la fila de encabezados"})
    resultado = []
    for n, fila in enumerate(filas, start=2):
        valores = {clave: (fila[j] if j < len(fila) else None) for clave, j in indice.items()}
        if all(v is None or str(v).strip() == "" for v in valores.values()):
            continue
        resultado.append({**valores, "_fila": n})
    return resultado, errores


def leer(contenido: bytes, tipo: str, hoja_csv: str | None) -> tuple[dict[str, list[dict]], list[dict]]:
    if not contenido:
        raise ArchivoInvalido("No llegó ningún archivo: elige la plantilla llena (Excel o CSV).")
    if len(contenido) > MAX_BYTES:
        raise ArchivoInvalido(f"El archivo pesa {len(contenido) / 1024 / 1024:.1f} MB y el máximo es 5 MB.")
    hojas: dict[str, list[dict]] = {h: [] for h in HOJAS}
    errores: list[dict] = []
    if tipo == TIPO_XLSX or contenido[:2] == b"PK":
        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise ArchivoInvalido("El archivo no es un Excel (.xlsx) que se pueda abrir. Descarga la plantilla y llénala de nuevo.")
        nombres = {_normalizar(n): n for n in libro.sheetnames}
        encontradas = 0
        for hoja in HOJAS:
            real = nombres.get(_normalizar(hoja))
            if real is None:
                continue
            encontradas += 1
            filas = [list(f) for f in libro[real].iter_rows(values_only=True)]
            if not filas:
                continue
            datos, errs = _filas_de(filas[0], filas[1:], hoja)
            hojas[hoja], errores = datos, errores + errs
        if not encontradas:
            raise ArchivoInvalido("El Excel no tiene ninguna de las hojas Insumos, Platos o Recetas: usa la plantilla del sistema.")
    elif tipo in TIPOS_CSV or tipo == "":
        if hoja_csv not in HOJAS:
            raise ArchivoInvalido("Con un CSV indica qué hoja es: Insumos, Platos o Recetas.")
        try:
            texto = contenido.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = contenido.decode("latin-1")
        separador = ";" if texto.split("\n", 1)[0].count(";") >= texto.split("\n", 1)[0].count(",") else ","
        filas = list(csv.reader(io.StringIO(texto), delimiter=separador))
        if filas:
            hojas[hoja_csv], errores = _filas_de(filas[0], filas[1:], hoja_csv)
    else:
        raise ArchivoInvalido(f"El archivo es de tipo «{tipo}»: sube la plantilla en Excel (.xlsx) o CSV.")
    total = sum(len(f) for f in hojas.values())
    if total > MAX_FILAS:
        raise ArchivoInvalido(f"El archivo tiene {total} filas y el máximo por carga es {MAX_FILAS}: divídelo en varias cargas.")
    return hojas, errores


# ---------------------------------------------------------------------------
# Validación y aplicación
# ---------------------------------------------------------------------------
MILES_RE = re.compile(r"^-?\d{1,3}(\.\d{3})+$")


def numero(valor) -> Decimal | None:
    if valor is None or (isinstance(valor, str) and valor.strip() == ""):
        return None
    if isinstance(valor, (int, float, Decimal)) and not isinstance(valor, bool):
        return Decimal(str(valor))
    texto = str(valor).strip().replace("$", "").replace(" ", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
    elif MILES_RE.match(texto):
        texto = texto.replace(".", "")
    try:
        return Decimal(texto)
    except InvalidOperation:
        raise ValueError(f"«{valor}» no es un número")


def texto(valor) -> str | None:
    if valor is None:
        return None
    t = str(valor).strip()
    if isinstance(valor, float) and valor.is_integer():
        t = str(int(valor))
    return t or None


def si_no(valor) -> bool | None:
    t = _normalizar(valor)
    if t == "":
        return None
    if t in ("si", "sí", "s", "x", "1", "true", "verdadero"):
        return True
    if t in ("no", "n", "0", "false", "falso"):
        return False
    raise ValueError(f"«{valor}» no es Sí o No")


SINONIMOS_UNIDAD = {
    "kg": "kg", "kilo": "kg", "kilos": "kg", "kilogramo": "kg", "kilogramos": "kg",
    "g": "g", "gr": "g", "grs": "g", "gramo": "g", "gramos": "g",
    "l": "L", "lt": "L", "lts": "L", "litro": "L", "litros": "L",
    "ml": "ml", "mililitro": "ml", "mililitros": "ml",
    "unidad": "unidad", "unidades": "unidad", "und": "unidad", "un": "unidad", "u": "unidad",
}


def unidad(valor) -> str:
    """"Kilos" → kg, "gr" → g, "Litros" → L, "und" → unidad. Lo que no se
    reconozca se deja tal cual y la validación dice la causa."""
    t = _normalizar(valor).rstrip(".")
    return SINONIMOS_UNIDAD.get(t, texto(valor) or "")


ENCABEZADO = {hoja: {c[1]: c[0] for c in cols} for hoja, cols in COLUMNAS.items()}


def _errores_pydantic(hoja: str, fila: int, exc: ValidationError) -> list[dict]:
    salida = []
    for e in exc.errors():
        clave = str(e["loc"][0]) if e["loc"] else ""
        mensaje = traducir_error(e)["mensaje"]
        salida.append({"hoja": hoja, "fila": fila, "columna": ENCABEZADO[hoja].get(clave, clave or "—"), "mensaje": mensaje})
    return salida


@dataclass
class Plan:
    insumos: list[dict] = field(default_factory=list)  # {datos, existencia, existente}
    platos: list[dict] = field(default_factory=list)  # {datos, categoria, existente}
    recetas: dict[str, list[dict]] = field(default_factory=dict)  # plato (minúsculas) → líneas
    errores: list[dict] = field(default_factory=list)
    resumen: dict[str, dict] = field(default_factory=dict)


def validar(db: Session, tenant_id: uuid.UUID, hojas: dict[str, list[dict]], errores_lectura: list[dict]) -> Plan:
    plan = Plan(errores=list(errores_lectura))
    insumos_db = {i.nombre.lower(): i for i in db.query(Insumo).filter(Insumo.tenant_id == tenant_id).all()}
    platos_db = {p.nombre.lower(): p for p in db.query(Producto).filter(Producto.tenant_id == tenant_id).all()}

    # --- Insumos -------------------------------------------------------------
    vistos: dict[str, int] = {}
    unidades_finales = {k: i.unidad for k, i in insumos_db.items()}
    nuevos = actualizados = 0
    for fila in hojas["Insumos"]:
        n = fila["_fila"]
        errores_fila = []
        try:
            minimo, costo, existencia = numero(fila.get("minimo")), numero(fila.get("costo_unitario")), numero(fila.get("existencia"))
        except ValueError as e:
            plan.errores.append({"hoja": "Insumos", "fila": n, "columna": "Mínimo / Costo / Existencia", "mensaje": str(e)})
            continue
        datos = {
            "nombre": texto(fila.get("nombre")) or "",
            "unidad": unidad(fila.get("unidad")),
            "grupo": texto(fila.get("grupo")),
            "minimo": minimo if minimo is not None else Decimal("0"),
            "costo_unitario": costo if costo is not None else Decimal("0"),
            "existencia_inicial": existencia if existencia is not None else Decimal("0"),
        }
        try:
            valido = si.InsumoCreate(**datos)
        except ValidationError as e:
            plan.errores.extend(_errores_pydantic("Insumos", n, e))
            continue
        clave = valido.nombre.lower()
        if clave in vistos:
            errores_fila.append({"hoja": "Insumos", "fila": n, "columna": "Nombre", "mensaje": f"{valido.nombre} ya aparece en la fila {vistos[clave]}"})
        existente = insumos_db.get(clave)
        if existente and existente.unidad != valido.unidad and not compatibles(existente.unidad, valido.unidad):
            errores_fila.append({
                "hoja": "Insumos", "fila": n, "columna": "Unidad",
                "mensaje": f"{existente.nombre} ya existe en {existente.unidad}: no se puede cargar en {valido.unidad}",
            })
        if errores_fila:
            plan.errores.extend(errores_fila)
            continue
        vistos[clave] = n
        unidades_finales[clave] = existente.unidad if existente else valido.unidad
        plan.insumos.append({"datos": valido, "existencia": existencia, "existente": existente})
        nuevos += existente is None
        actualizados += existente is not None
    plan.resumen["Insumos"] = {"filas": len(hojas["Insumos"]), "nuevos": nuevos, "actualizados": actualizados}

    # --- Platos --------------------------------------------------------------
    vistos = {}
    nombres_finales = set(platos_db)
    nuevos = actualizados = 0
    for fila in hojas["Platos"]:
        n = fila["_fila"]
        try:
            precio = numero(fila.get("precio"))
            cocina = si_no(fila.get("va_a_cocina"))
        except ValueError as e:
            plan.errores.append({"hoja": "Platos", "fila": n, "columna": "Precio / Va a cocina", "mensaje": str(e)})
            continue
        categoria = texto(fila.get("categoria"))
        try:
            valido = sc.ProductoCreate(
                nombre=texto(fila.get("nombre")) or "",
                precio=precio if precio is not None else Decimal("0"),
                va_a_cocina=True if cocina is None else cocina,
                sku=texto(fila.get("sku")),
                descripcion=texto(fila.get("descripcion")),
            )
            if categoria:
                sc.CategoriaCreate(nombre=categoria)
        except ValidationError as e:
            plan.errores.extend(_errores_pydantic("Platos", n, e))
            continue
        clave = valido.nombre.lower()
        if clave in vistos:
            plan.errores.append({"hoja": "Platos", "fila": n, "columna": "Nombre", "mensaje": f"{valido.nombre} ya aparece en la fila {vistos[clave]}"})
            continue
        existente = platos_db.get(clave)
        if existente and existente.tipo == "combo":
            plan.errores.append({"hoja": "Platos", "fila": n, "columna": "Nombre", "mensaje": f"{existente.nombre} es un combo: los combos se editan en Carta → Combos"})
            continue
        vistos[clave] = n
        nombres_finales.add(clave)
        plan.platos.append({"datos": valido, "categoria": categoria, "existente": existente})
        nuevos += existente is None
        actualizados += existente is not None
    plan.resumen["Platos"] = {"filas": len(hojas["Platos"]), "nuevos": nuevos, "actualizados": actualizados}

    # --- Recetas -------------------------------------------------------------
    for fila in hojas["Recetas"]:
        n = fila["_fila"]
        plato, insumo = texto(fila.get("plato")) or "", texto(fila.get("insumo")) or ""
        errores_fila = []
        if plato.lower() not in nombres_finales:
            errores_fila.append({"hoja": "Recetas", "fila": n, "columna": "Plato", "mensaje": f"El plato «{plato}» no está en la hoja Platos ni en la carta"})
        elif plato.lower() in platos_db and platos_db[plato.lower()].tipo == "combo":
            errores_fila.append({"hoja": "Recetas", "fila": n, "columna": "Plato", "mensaje": f"«{plato}» es un combo: no lleva receta propia"})
        if insumo.lower() not in unidades_finales:
            errores_fila.append({"hoja": "Recetas", "fila": n, "columna": "Insumo", "mensaje": f"El insumo «{insumo}» no está en la hoja Insumos ni en el inventario"})
        try:
            cantidad, merma = numero(fila.get("cantidad")), numero(fila.get("merma_porcentaje"))
            linea = sc.RecetaItemIn(
                insumo_id=uuid.uuid4(), cantidad=cantidad if cantidad is not None else Decimal("0"),
                unidad=unidad(fila.get("unidad")), merma_porcentaje=merma if merma is not None else Decimal("0"),
            )
        except ValueError as e:
            if isinstance(e, ValidationError):
                errores_fila.extend(_errores_pydantic("Recetas", n, e))
            else:
                errores_fila.append({"hoja": "Recetas", "fila": n, "columna": "Cantidad / Merma %", "mensaje": str(e)})
            plan.errores.extend(errores_fila)
            continue
        unidad_insumo = unidades_finales.get(insumo.lower())
        if unidad_insumo and not compatibles(linea.unidad, unidad_insumo):
            errores_fila.append({
                "hoja": "Recetas", "fila": n, "columna": "Unidad",
                "mensaje": f"{insumo} se lleva en {unidad_insumo}: la receta no puede ir en {linea.unidad}",
            })
        lineas = plan.recetas.setdefault(plato.lower(), [])
        if any(l["insumo"] == insumo.lower() for l in lineas):
            errores_fila.append({"hoja": "Recetas", "fila": n, "columna": "Insumo", "mensaje": f"{insumo} ya está en la receta de {plato}: súmalo en una sola fila"})
        if errores_fila:
            plan.errores.extend(errores_fila)
            continue
        lineas.append({"insumo": insumo.lower(), "linea": linea})
    plan.resumen["Recetas"] = {"filas": len(hojas["Recetas"]), "platos": len(plan.recetas)}
    plan.errores.sort(key=lambda e: (HOJAS.index(e["hoja"]), e["fila"]))
    return plan


def informe(plan: Plan, aplicado: bool = False) -> dict:
    return {
        "valido": not plan.errores,
        "aplicado": aplicado,
        "total_errores": len(plan.errores),
        "errores": plan.errores[:500],
        "hojas": [{"hoja": hoja, **plan.resumen.get(hoja, {"filas": 0})} for hoja in HOJAS],
    }


def aplicar(db: Session, tenant_id: uuid.UUID, plan: Plan, usuario_nombre: str) -> None:
    """Todo o nada: quien llama hace commit al final (o rollback si algo falla)."""
    insumos = {i.nombre.lower(): i for i in db.query(Insumo).filter(Insumo.tenant_id == tenant_id).all()}
    for item in plan.insumos:
        d = item["datos"]
        insumo = item["existente"]
        nuevo = insumo is None
        if nuevo:
            insumo = Insumo(tenant_id=tenant_id, nombre=d.nombre, unidad=d.unidad, existencia=Decimal("0"))
            db.add(insumo)
        insumo.grupo = d.grupo if d.grupo else insumo.grupo
        insumo.minimo = d.minimo
        if d.costo_unitario:
            insumo.costo_unitario = d.costo_unitario
        db.flush()
        insumos[insumo.nombre.lower()] = insumo
        if item["existencia"] is not None:
            meta = convertir(item["existencia"], d.unidad, insumo.unidad)
            diferencia = Decimal(meta) - Decimal(insumo.existencia)
            if diferencia != 0:
                services.registrar(
                    db, tenant_id, insumo, "inicial" if nuevo else "ajuste", diferencia,
                    motivo="Carga inicial", usuario_nombre=usuario_nombre, permitir_negativo=True,
                )

    categorias = {c.nombre.lower(): c for c in db.query(Categoria).filter(Categoria.tenant_id == tenant_id).all()}
    platos = {p.nombre.lower(): p for p in db.query(Producto).filter(Producto.tenant_id == tenant_id).all()}
    for item in plan.platos:
        d = item["datos"]
        categoria = None
        if item["categoria"]:
            categoria = categorias.get(item["categoria"].lower())
            if categoria is None:
                categoria = Categoria(tenant_id=tenant_id, nombre=item["categoria"])
                db.add(categoria)
                db.flush()
                categorias[categoria.nombre.lower()] = categoria
        plato = item["existente"]
        if plato is None:
            plato = Producto(tenant_id=tenant_id, tipo="plato", nombre=d.nombre)
            db.add(plato)
        plato.precio = d.precio
        plato.va_a_cocina = d.va_a_cocina
        if d.sku:
            plato.sku = d.sku
        if d.descripcion:
            plato.descripcion = d.descripcion
        if categoria is not None:
            plato.categoria_id = categoria.id
        db.flush()
        platos[plato.nombre.lower()] = plato

    for nombre_plato, lineas in plan.recetas.items():
        plato = platos[nombre_plato]
        db.query(RecetaItem).filter(RecetaItem.producto_id == plato.id).delete(synchronize_session=False)
        for i, l in enumerate(lineas):
            linea = l["linea"]
            db.add(RecetaItem(
                tenant_id=tenant_id, producto_id=plato.id, insumo_id=insumos[l["insumo"]].id,
                cantidad=linea.cantidad, unidad=linea.unidad, merma_porcentaje=linea.merma_porcentaje, orden=i,
            ))
    db.flush()
