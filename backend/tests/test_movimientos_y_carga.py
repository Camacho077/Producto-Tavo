"""
FASE 3 (2026-10-12): mermas con motivo, ajuste por conteo físico, historial
general de movimientos y carga inicial desde la plantilla (Excel o CSV),
revisada fila por fila antes de guardar nada.
"""
import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import API_V1 as API, comprar, crear_insumo, crear_plato, existencia

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def s(client, tenant_factory, usuario_factory, auth_headers_factory):
    tenant = tenant_factory()
    return client, auth_headers_factory(usuario_factory(tenant, nombre_completo="Ana Bodega"))


# ---------------------------------------------------------------------------
# Mermas
# ---------------------------------------------------------------------------
def test_la_merma_resta_con_motivo_quien_y_valor(s):
    client, h = s
    queso = crear_insumo(client, h, "Queso", "kg", costo="20000", existencia="2")

    r = client.post(f"{API}/inventario/insumos/{queso['id']}/merma", headers=h, json={"cantidad": "300", "unidad": "g", "motivo": "vencido"})

    assert r.status_code == 201, r.text
    m = r.json()
    assert m["tipo"] == "merma" and Decimal(m["cantidad"]) == Decimal("-0.3")
    assert m["motivo"] == "Vencido" and m["usuario_nombre"] == "Ana Bodega"
    assert Decimal(m["valor"]) == Decimal("-6000")
    assert existencia(client, h, queso["id"]) == Decimal("1.7")


def test_merma_otro_exige_detalle_y_lo_guarda(s):
    client, h = s
    leche = crear_insumo(client, h, "Leche", "L", existencia="5")
    sin = client.post(f"{API}/inventario/insumos/{leche['id']}/merma", headers=h, json={"cantidad": "1", "unidad": "L", "motivo": "otro"})
    assert sin.status_code == 422 and "detalle es obligatorio" in sin.text

    con = client.post(f"{API}/inventario/insumos/{leche['id']}/merma", headers=h, json={"cantidad": "1", "unidad": "L", "motivo": "otro", "detalle": "Se cortó  con el calor"})
    assert con.status_code == 201 and con.json()["motivo"] == "Otro: Se cortó con el calor"


def test_merma_no_deja_existencia_negativa_y_dice_cuanto_queda(s):
    client, h = s
    arroz = crear_insumo(client, h, "Arroz", "kg", existencia="1")
    r = client.post(f"{API}/inventario/insumos/{arroz['id']}/merma", headers=h, json={"cantidad": "1500", "unidad": "g", "motivo": "danado"})
    assert r.status_code == 409
    assert "quedan 1 kg" in r.json()["detail"] and "1,5 kg" in r.json()["detail"] and "conteo físico" in r.json()["detail"]
    assert existencia(client, h, arroz["id"]) == 1


def test_merma_en_unidad_de_otra_familia_se_rechaza(s):
    client, h = s
    huevo = crear_insumo(client, h, "Huevo", "unidad", existencia="30")
    r = client.post(f"{API}/inventario/insumos/{huevo['id']}/merma", headers=h, json={"cantidad": "1", "unidad": "kg", "motivo": "danado"})
    assert r.status_code == 409 and "se lleva en unidad" in r.json()["detail"]


# ---------------------------------------------------------------------------
# Conteo físico
# ---------------------------------------------------------------------------
def test_el_conteo_ajusta_a_lo_contado_y_valoriza_la_diferencia(s):
    client, h = s
    arroz = crear_insumo(client, h, "Arroz", "kg", costo="4000", existencia="10")
    sal = crear_insumo(client, h, "Sal", "kg", costo="1000", existencia="2")
    aceite = crear_insumo(client, h, "Aceite", "L", costo="9000", existencia="3")

    r = client.post(f"{API}/inventario/conteos", headers=h, json={"motivo": "Cierre de mes", "lineas": [
        {"insumo_id": arroz["id"], "contado": "8.5"},
        {"insumo_id": sal["id"], "contado": "2"},
        {"insumo_id": aceite["id"], "contado": "4"},
    ]})

    assert r.status_code == 201, r.text
    c = r.json()
    assert c["codigo"] == "CON-000001" and c["lineas"] == 3 and c["con_diferencia"] == 2
    # -1,5 kg × 4.000 + 1 L × 9.000 = 3.000
    assert Decimal(c["valor_diferencia"]) == Decimal("3000")
    assert existencia(client, h, arroz["id"]) == Decimal("8.5")
    assert existencia(client, h, aceite["id"]) == 4

    movs = client.get(f"{API}/inventario/movimientos", headers=h, params={"tipo": "ajuste"}).json()["items"]
    assert {m["insumo_nombre"] for m in movs} == {"Arroz", "Aceite"}  # sin diferencia, sin movimiento
    assert all(m["documento_codigo"] == "CON-000001" and m["motivo"] == "Cierre de mes" for m in movs)

    detalle = client.get(f"{API}/inventario/conteos/{c['id']}", headers=h).json()
    assert [Decimal(i["diferencia"]) for i in detalle["items"]] == [Decimal("-1.5"), 0, 1]
    assert client.post(f"{API}/inventario/conteos", headers=h, json={"lineas": [{"insumo_id": sal["id"], "contado": "1"}]}).json()["codigo"] == "CON-000002"


def test_el_conteo_no_repite_insumos(s):
    client, h = s
    sal = crear_insumo(client, h, "Sal", "kg")
    r = client.post(f"{API}/inventario/conteos", headers=h, json={"lineas": [{"insumo_id": sal["id"], "contado": "1"}, {"insumo_id": sal["id"], "contado": "2"}]})
    assert r.status_code == 409 and "repite un insumo" in r.json()["detail"]


def test_conteo_con_insumo_de_otra_empresa_da_404_y_no_toca_nada(client, tenant_factory, usuario_factory, auth_headers_factory):
    ha = auth_headers_factory(usuario_factory(tenant_factory()))
    hb = auth_headers_factory(usuario_factory(tenant_factory()))
    ajeno = crear_insumo(client, ha, "Sal", "kg", existencia="5")
    r = client.post(f"{API}/inventario/conteos", headers=hb, json={"lineas": [{"insumo_id": ajeno["id"], "contado": "0"}]})
    assert r.status_code == 404
    assert existencia(client, ha, ajeno["id"]) == 5


# ---------------------------------------------------------------------------
# Historial general
# ---------------------------------------------------------------------------
def test_historial_general_filtra_por_tipo_insumo_y_busca(s):
    client, h = s
    arroz = crear_insumo(client, h, "Arroz", "kg", costo="4000")
    sal = crear_insumo(client, h, "Sal", "kg", existencia="3")
    comprar(client, h, [(arroz["id"], "10", "kg", "4000")])
    client.post(f"{API}/inventario/insumos/{sal['id']}/merma", headers=h, json={"cantidad": "1", "unidad": "kg", "motivo": "derramado"})

    todos = client.get(f"{API}/inventario/movimientos", headers=h).json()
    assert todos["total"] == 3
    assert [m["tipo"] for m in client.get(f"{API}/inventario/movimientos", headers=h, params={"tipo": "merma"}).json()["items"]] == ["merma"]
    assert {m["tipo"] for m in client.get(f"{API}/inventario/movimientos", headers=h, params={"insumo_id": arroz["id"]}).json()["items"]} == {"compra"}
    assert client.get(f"{API}/inventario/movimientos", headers=h, params={"q": "derram"}).json()["total"] == 1
    assert client.get(f"{API}/inventario/movimientos", headers=h, params={"desde": "2000-01-01", "hasta": "2000-01-02"}).json()["total"] == 0

    csv = client.get(f"{API}/inventario/movimientos", headers=h, params={"formato": "csv"})
    assert csv.status_code == 200 and "Merma" in csv.text and "Se cayó o se derramó" in csv.text


def test_historial_sin_costos_no_muestra_valores(client, tenant_factory, rol_factory, usuario_factory, auth_headers_factory):
    tenant = tenant_factory()
    ha = auth_headers_factory(usuario_factory(tenant))
    crear_insumo(client, ha, "Arroz", "kg", costo="4000", existencia="2")
    bodega = rol_factory(tenant, nombre="Bodega", permisos={"movimientos": ["ver"]})
    hb = auth_headers_factory(usuario_factory(tenant, rol=bodega))
    item = client.get(f"{API}/inventario/movimientos", headers=hb).json()["items"][0]
    assert item["valor"] is None and item["costo_unitario"] is None
    assert "Valor a costo" not in client.get(f"{API}/inventario/movimientos", headers=hb, params={"formato": "csv"}).text


# ---------------------------------------------------------------------------
# Carga inicial
# ---------------------------------------------------------------------------
def _excel(insumos=(), platos=(), recetas=()):
    libro = Workbook()
    libro.remove(libro.active)
    for nombre, encabezados, filas in (
        ("Insumos", ["Nombre *", "Unidad *", "Grupo", "Mínimo", "Costo unitario", "Existencia"], insumos),
        ("Platos", ["Nombre *", "Categoría", "Precio *", "Va a cocina", "Código", "Descripción"], platos),
        ("Recetas", ["Plato *", "Insumo *", "Cantidad *", "Unidad *", "Merma %"], recetas),
    ):
        hoja = libro.create_sheet(nombre)
        hoja.append(encabezados)
        for fila in filas:
            hoja.append(list(fila))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _subir(client, h, accion, contenido, tipo=TIPO_XLSX, hoja=None):
    return client.post(
        f"{API}/inventario/carga-inicial/{accion}", headers={**h, "Content-Type": tipo},
        params={"hoja": hoja} if hoja else None, content=contenido,
    )


def test_la_plantilla_excel_trae_las_hojas_y_los_encabezados(s):
    client, h = s
    r = client.get(f"{API}/inventario/carga-inicial/plantilla", headers=h)
    assert r.status_code == 200 and r.headers["content-type"] == TIPO_XLSX
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Instrucciones", "Insumos", "Platos", "Recetas"]
    assert [c.value for c in libro["Insumos"][1]][:2] == ["Nombre *", "Unidad *"]


def test_la_plantilla_csv_es_por_hoja(s):
    client, h = s
    r = client.get(f"{API}/inventario/carga-inicial/plantilla", headers=h, params={"formato": "csv", "hoja": "Recetas"})
    assert r.status_code == 200 and r.content.decode("utf-8-sig").startswith("Plato *;Insumo *")
    assert client.get(f"{API}/inventario/carga-inicial/plantilla", headers=h, params={"formato": "csv"}).status_code == 409


def test_validar_informa_cada_error_con_hoja_fila_y_columna_sin_guardar_nada(s):
    client, h = s
    archivo = _excel(
        insumos=[("Arroz", "kilos", "Abarrotes", 5, 4200, 25), ("Pollo", "libras", None, None, None, None), ("Arroz", "kg", None, None, None, None)],
        platos=[("Arroz con pollo", "Almuerzos", 14000, "Sí", None, None), ("Jugo", None, "caro", None, None, None)],
        recetas=[("Arroz con pollo", "Arroz", 200, "g", 0), ("Arroz con pollo", "Papa", 100, "g", 0), ("Bandeja", "Arroz", 1, "kg", 0)],
    )
    r = _subir(client, h, "validar", archivo)

    assert r.status_code == 200, r.text
    informe = r.json()
    assert informe["valido"] is False and informe["aplicado"] is False
    errores = {(e["hoja"], e["fila"], e["columna"]) for e in informe["errores"]}
    assert ("Insumos", 3, "Unidad") in errores  # "libras" no es una unidad
    assert ("Insumos", 4, "Nombre") in errores  # Arroz repetido
    assert ("Platos", 3, "Precio / Va a cocina") in errores
    assert ("Recetas", 3, "Insumo") in errores and ("Recetas", 4, "Plato") in errores
    assert any("Papa" in e["mensaje"] for e in informe["errores"])
    unidad = next(e for e in informe["errores"] if (e["fila"], e["columna"]) == (3, "Unidad"))
    assert unidad["mensaje"] == "Debe ser una de estas opciones: kg, g, L, ml o unidad"  # en español, nunca el de Pydantic
    # Nada se guardó.
    assert client.get(f"{API}/inventario/insumos", headers=h).json()["total"] == 0


def test_aplicar_con_errores_responde_422_con_el_informe_y_no_guarda(s):
    client, h = s
    r = _subir(client, h, "aplicar", _excel(insumos=[("Arroz", "kg", None, None, None, 10), ("Sal", "litros de sal", None, None, None, None)]))
    assert r.status_code == 422
    assert "No se guardó nada" in r.json()["detail"] and r.json()["datos"]["total_errores"] == 1
    assert client.get(f"{API}/inventario/insumos", headers=h).json()["total"] == 0


def test_aplicar_crea_insumos_platos_y_recetas_con_existencia_como_movimiento(s):
    client, h = s
    archivo = _excel(
        insumos=[("Arroz", "kilos", "Abarrotes", 5, 4200, 25), ("Pollo", "kg", "Carnes", 2, "18.000", "3,5")],
        platos=[("Arroz con pollo", "Almuerzos", 14000, "Sí", None, "Con ensalada"), ("Gaseosa", "Bebidas", 3500, "No", None, None)],
        recetas=[("Arroz con pollo", "Arroz", 200, "g", 0), ("Arroz con pollo", "Pollo", 150, "g", 10)],
    )
    validado = _subir(client, h, "validar", archivo).json()
    assert validado["valido"] is True
    assert validado["hojas"][0] == {"hoja": "Insumos", "filas": 2, "nuevos": 2, "actualizados": 0}

    r = _subir(client, h, "aplicar", archivo)
    assert r.status_code == 200 and r.json()["aplicado"] is True

    insumos = {i["nombre"]: i for i in client.get(f"{API}/inventario/insumos", headers=h).json()["items"]}
    assert Decimal(insumos["Pollo"]["existencia"]) == Decimal("3.5") and Decimal(insumos["Pollo"]["costo_unitario"]) == 18000
    movs = client.get(f"{API}/inventario/movimientos", headers=h).json()["items"]
    assert {(m["tipo"], m["motivo"]) for m in movs} == {("inicial", "Carga inicial")}

    platos = {p["nombre"]: p for p in client.get(f"{API}/catalogo/productos", headers=h).json()["items"]}
    assert platos["Gaseosa"]["va_a_cocina"] is False
    plato = client.get(f"{API}/catalogo/productos/{platos['Arroz con pollo']['id']}", headers=h).json()
    assert len(plato["receta"]) == 2 and plato["categoria_nombre"] == "Almuerzos"


def test_volver_a_cargar_actualiza_y_ajusta_la_diferencia(s):
    client, h = s
    arroz = crear_insumo(client, h, "Arroz", "kg", existencia="10")
    sal = crear_insumo(client, h, "Sal", "kg")
    plato = crear_plato(client, h, "Arroz blanco", receta=[(arroz["id"], "100", "g"), (sal["id"], "2", "g")])

    archivo = _excel(insumos=[("arroz", "kg", None, 3, None, 12)], recetas=[("Arroz blanco", "Arroz", 150, "g", 0)])
    informe = _subir(client, h, "aplicar", archivo).json()
    assert informe["hojas"][0]["actualizados"] == 1

    assert existencia(client, h, arroz["id"]) == 12
    ajuste = client.get(f"{API}/inventario/movimientos", headers=h, params={"tipo": "ajuste"}).json()["items"]
    assert len(ajuste) == 1 and Decimal(ajuste[0]["cantidad"]) == 2
    receta = client.get(f"{API}/catalogo/productos/{plato['id']}", headers=h).json()["receta"]
    assert [(r["insumo_id"], Decimal(r["cantidad"])) for r in receta] == [(arroz["id"], 150)]  # se reemplaza entera


def test_carga_csv_por_hoja_con_coma_decimal(s):
    client, h = s
    csv = "Nombre *;Unidad *;Grupo;Mínimo;Costo unitario;Existencia\r\nAceite;litros;;1;9.500;2,5\r\n".encode("utf-8-sig")
    r = _subir(client, h, "aplicar", csv, tipo="text/csv", hoja="Insumos")
    assert r.status_code == 200, r.text
    aceite = client.get(f"{API}/inventario/insumos", headers=h).json()["items"][0]
    assert aceite["unidad"] == "L" and Decimal(aceite["existencia"]) == Decimal("2.5") and Decimal(aceite["costo_unitario"]) == 9500


def test_archivo_ilegible_o_vacio_dice_la_causa(s):
    client, h = s
    assert "No llegó ningún archivo" in _subir(client, h, "validar", b"").json()["detail"]
    assert "no es un Excel" in _subir(client, h, "validar", b"PK basura").json()["detail"]
    assert "indica qué hoja" in _subir(client, h, "validar", b"a;b", tipo="text/csv").json()["detail"]
    assert "image/png" in _subir(client, h, "validar", b"\x89PNG", tipo="image/png").json()["detail"]


def test_rol_sin_permiso_de_carta_no_puede_cargar_platos(client, tenant_factory, rol_factory, usuario_factory, auth_headers_factory):
    tenant = tenant_factory()
    bodega = rol_factory(tenant, nombre="Bodega", permisos={"insumos": ["ver", "crear"]})
    h = auth_headers_factory(usuario_factory(tenant, rol=bodega))
    r = _subir(client, h, "validar", _excel(platos=[("Jugo", None, 3000, None, None, None)]))
    assert r.status_code == 403 and "solo la hoja de insumos" in r.json()["detail"]
    assert _subir(client, h, "aplicar", _excel(insumos=[("Sal", "kg", None, None, None, None)])).status_code == 200
